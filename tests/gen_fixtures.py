# Генератор тестових файлів і конфігурації етапу C (регресійні тести Power Query Engine)
import os, sys, datetime as dt
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
OUT = sys.argv[1]; FX = os.path.join(OUT, 'fixtures'); os.makedirs(FX, exist_ok=True)

HDR = ['Код', 'Назва', 'Кількість', 'Ціна, грн.', 'Дата', 'Примітка']
GOOD = [
    ['A001', 'Болт  М6', '10', '1,50', '01.10.2026', ''],
    ['A002', 'Гайка М6', '20', '0,75', '02.10.2026', 'прим'],
    ['A003', 'Шайба', '100', '1 234,50', '03.10.2026', ''],
    ['A004', 'Кутник', '3', '12', '04.10.2026', ''],
    ['A005', 'Дюбель', '7', '0,2', '05.10.2026', ''],
]
TEN = GOOD + [
    ['A006', 'Ковпак', 'abc', '2', '06.10.2026', ''],
    ['A007', 'Скоба', '5', '3,1', '07.10.2026', ''],
    ['A008', 'Хомут', '1,2,3', '4', '08.10.2026', ''],
    ['A009', 'Шуруп', '50', '0,1', '09.10.2026', ''],
    ['A010', 'Цвях', '500', '0,05', '10.10.2026', ''],
]
def csv(name, header, rows, enc='utf-8-sig'):
    txt = '\r\n'.join(';'.join(r) for r in [header] + rows) + '\r\n'
    open(os.path.join(FX, name), 'w', encoding=enc, newline='').write(txt)

csv('C01_ok.csv', HDR, GOOD)
csv('C02_service_rows.csv', HDR, GOOD[:3] + [['Всього', '', '130', '', '', '']] + GOOD[3:] +
    [['A006', 'Ковпак', '1', 'х', '06.10.2026', ''], ['Разом', '', '', '', '', '']])
csv('C03_type_errors.csv', HDR, TEN)
csv('C04_threshold.csv', HDR, TEN)
csv('C05_strict.csv', HDR, TEN)
csv('C06_lenient.csv', HDR, TEN)
csv('C07_missing_required.csv', [h for h in HDR if h != 'Назва'], [[v for i, v in enumerate(r) if i != 1] for r in GOOD])
csv('C08_missing_optional.csv', [h for h in HDR if h != 'Дата'], [[v for i, v in enumerate(r) if i != 4] for r in GOOD])
csv('C09_dup_header.csv', ['Код', 'Назва', 'Кількість', 'Ціна, грн.', 'Дата', 'Назва'],
    [r[:5] + ['ХИБНА ' + r[0]] for r in GOOD])
csv('C10_empty.csv', HDR, [])
csv('C13_utf8_no_bom.csv', HDR, GOOD, enc='utf-8')
csv('C14_cp1251.csv', HDR, GOOD, enc='cp1251')
csv('C15_disabled.csv', HDR, GOOD)
csv('C16_dup_period.csv', HDR, GOOD)
csv('N01_no_schema.csv', HDR, GOOD)

def xrows(rows):
    out = []
    for r in rows:
        d = dt.datetime.strptime(r[4], '%d.%m.%Y') if r[4] else None
        num = lambda s: float(s.replace(' ', '').replace(',', '.')) if s not in ('', 'х') else s
        out.append([r[0], r[1], int(r[2]) if r[2].isdigit() else r[2], num(r[3]), d, r[5] or None])
    return out
def sheet(ws, header, rows):
    ws.append(header)
    for r in rows: ws.append(r)

wb = Workbook(); ws = wb.active; ws.title = 'Дані'
sheet(ws, HDR, xrows(GOOD[:2]) + [['Всього', None, 30, None, None, None]] + xrows(GOOD[2:]))
t = Table(displayName='tblData', ref=f'A1:F{ws.max_row}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9')
ws.add_table(t); wb.save(os.path.join(FX, 'X01_table.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = 'Аркуш1'
rows = xrows(GOOD); rows[2][2] = '#N/A'     # клітинка-помилка #N/A у колонці «Кількість»
sheet(ws, HDR, rows); wb.save(os.path.join(FX, 'X02_na_error.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = 'Дані'
sheet(ws, HDR, xrows(GOOD)); wb.save(os.path.join(FX, 'X03_wrong_sheet.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = 'Дані1'; sheet(ws, HDR, xrows(GOOD))
ws2 = wb.create_sheet('Дані2'); sheet(ws2, HDR, xrows(GOOD)); wb.save(os.path.join(FX, 'X04_ambiguous.xlsx'))

open(os.path.join(FX, 'X05_corrupt.xlsx'), 'w', encoding='utf-8').write('це не книга Excel\r\n')
open(os.path.join(FX, 'X06_unsupported.json'), 'w', encoding='utf-8').write('{"a": 1}\r\n')
print('fixtures:', len(os.listdir(FX)))

# ---------------- конфігурація тестової книги ----------------
# реєстр: SourceID, TargetType, FullPath, Path, FileName, TableSheetName, PeriodType, PeriodNumber, IsRequired, Enabled, ErrorPolicy, ErrorThreshold
REG = [
    ['C01', 'TestCsv', 'C01_ok.csv', None, 'Місяць', 1, None, None, None, None, 'Еталонний файл'],
    ['C02', 'TestCsv', 'C02_service_rows.csv', None, 'Місяць', 2, None, None, None, None, 'Службові рядки Всього, Разом, х'],
    ['C03', 'TestCsv', 'C03_type_errors.csv', None, 'Місяць', 3, None, None, None, None, '2 з 10 рядків з помилкою типу, Report'],
    ['C04', 'TestCsv', 'C04_threshold.csv', None, 'Місяць', 4, None, None, None, 0.1, 'Те саме, поріг 10% (факт 20%)'],
    ['C05', 'TestCsv', 'C05_strict.csv', None, 'Місяць', 5, None, None, 'Strict', None, 'Те саме, Strict'],
    ['C06', 'TestCsv', 'C06_lenient.csv', None, 'Місяць', 6, None, None, 'Lenient', None, 'Те саме, Lenient'],
    ['C07', 'TestCsv', 'C07_missing_required.csv', None, 'Місяць', 7, None, None, None, None, 'Немає обов\'язкової колонки Назва'],
    ['C08', 'TestCsv', 'C08_missing_optional.csv', None, 'Місяць', 8, None, None, None, None, 'Немає необов\'язкової колонки Дата'],
    ['C09', 'TestCsv', 'C09_dup_header.csv', None, 'Місяць', 9, None, None, None, None, 'Колонка Назва двічі (друга хибна)'],
    ['C10', 'TestCsv', 'C10_empty.csv', None, 'Місяць', 10, None, None, None, None, 'Лише заголовок'],
    ['C11', 'TestCsv', 'C11_missing.csv', None, 'Місяць', 11, False, None, None, None, 'Файла немає, необов\'язковий'],
    ['C12', 'TestCsv', 'C12_missing.csv', None, 'Місяць', 12, True, None, None, None, 'Файла немає, обов\'язковий'],
    ['C13', 'TestCsv', 'C13_utf8_no_bom.csv', None, 'Квартал', 1, None, None, None, None, 'UTF-8 без BOM при CsvEncoding 1251'],
    ['C14', 'TestCsv', 'C14_cp1251.csv', None, 'Квартал', 2, None, None, None, None, 'Кодування 1251'],
    ['C15', 'TestCsv', 'C15_disabled.csv', None, 'Квартал', 4, None, False, None, None, 'Enabled = FALSE'],
    ['C16', 'TestCsv', 'C16_dup_period.csv', None, 'Місяць', 1, None, None, None, None, 'Той самий період, що C01'],
    ['C17', 'TestCsv', 'C01_ok.csv', None, 'Квартал', 3, None, None, None, None, 'Той самий файл, що C01'],
    ['X01', 'TestXlsx', 'X01_table.xlsx', 'tblData', 'Місяць', 1, None, None, None, None, 'Таблиця за іменем, рядок Всього'],
    ['X02', 'TestXlsx', 'X02_na_error.xlsx', 'Аркуш1', 'Місяць', 2, None, None, None, None, 'Аркуш, клітинка #N/A'],
    ['X03', 'TestXlsx', 'X03_wrong_sheet.xlsx', 'НемаТакого', 'Місяць', 3, None, None, None, None, 'Неіснуючий TableSheetName'],
    ['X04', 'TestXlsx', 'X04_ambiguous.xlsx', None, 'Місяць', 4, None, None, None, None, 'Два аркуші, TableSheetName порожнє'],
    ['X05', 'TestXlsx', 'X05_corrupt.xlsx', None, 'Місяць', 5, None, None, None, None, 'Пошкоджений файл'],
    ['X06', 'TestXlsx', 'X06_unsupported.json', None, 'Місяць', 6, None, None, None, None, 'Непідтримуване розширення'],
    ['N01', 'TestNoSchema', 'N01_no_schema.csv', None, 'Рік', 1, None, None, None, None, 'Призначення без схеми'],
]
SCHEMA = []
for tt in ('TestCsv', 'TestXlsx'):
    SCHEMA += [
        [tt, 'Код', 'Code', 'text', True, True, None, None],
        [tt, 'Назва', 'Name', 'text', True, True, 'CleanSpaces', None],
        [tt, 'Кількість', 'Qty', 'number', True, None, None, None],
        [tt, 'Ціна, грн.', 'Price', 'number', False, None, 'NumText', None],
        [tt, 'Дата', 'Date', 'date', False, None, None, None],
        [tt, 'Примітка', None, 'text', False, None, None, False],
    ]
FILTERS = [
    ['TestCsv', 'Код', 'StartsWith', 'Всього|Разом', True],
    ['TestCsv', 'Ціна, грн.', 'Equals', 'х|x', True],
    ['TestXlsx', 'Код', 'StartsWith', 'Всього', True],
]
TARGETS = [['TestCsv', 'Тести Engine: CSV'], ['TestXlsx', 'Тести Engine: Excel'], ['TestNoSchema', 'Тести Engine: призначення без схеми']]

# очікування зведення: порожньо = не перевіряти, "-" = очікується порожнє значення
EXP_COLS = ['SourceID', 'Status', 'QualityFlag', 'Rows', 'RowsDropped', 'ErrorRows', 'MissingRequired', 'MissingOptional', 'DuplicateHeaders', 'Note']
EXP = [
    ['C01', 'OK', 'OK', 5, 0, 0, '-', '-', '-', 'Еталон'],
    ['C02', 'OK', 'OK', 5, 3, 0, None, None, None, 'Відсічено Всього, х, Разом'],
    ['C03', 'ERROR', 'WARN', 10, 0, 2, None, None, None, 'Report без порогу'],
    ['C04', 'ERROR', 'FAIL', 10, None, 2, None, None, None, 'Поріг 10% < 20%'],
    ['C05', 'ERROR', 'FAIL', 10, None, 2, None, None, None, 'Strict'],
    ['C06', 'ERROR', 'WARN', 10, None, 2, None, None, None, 'Lenient: звіт бачить помилки до заміни на null'],
    ['C07', 'ERROR', 'FAIL', None, None, None, 'Назва', None, None, None],
    ['C08', 'WARN', 'OK', 5, None, 0, '-', 'Дата', None, None],
    ['C09', 'WARN', 'OK', 5, None, 0, None, None, 'Назва', 'Тихий дубль колонки схеми'],
    ['C10', 'EMPTY', 'OK', 0, None, None, None, None, None, None],
    ['C11', 'NODATA', '-', None, None, None, None, None, None, 'Необов\'язковий файл відсутній'],
    ['C12', 'FILE_ERROR', 'FAIL', None, None, None, None, None, None, None],
    ['C13', 'ERROR', 'FAIL', None, None, None, 'Код, Назва, Кількість', None, None, 'Спотворені заголовки'],
    ['C14', 'OK', 'OK', 5, None, 0, None, None, None, None],
    ['C15', 'ABSENT', None, None, None, None, None, None, None, 'Вимкнене джерело відсутнє у зведенні'],
    ['C16', 'OK', 'OK', 5, None, None, None, None, None, None],
    ['C17', 'OK', 'OK', 5, None, None, None, None, None, None],
    ['X01', 'OK', 'OK', 5, 1, 0, None, None, None, None],
    ['X02', 'ERROR', 'WARN', 5, None, 1, None, None, None, '#N/A у джерелі'],
    ['X03', 'READ_ERROR', 'FAIL', None, None, None, None, None, None, None],
    ['X04', 'READ_ERROR', 'FAIL', None, None, None, None, None, None, None],
    ['X05', 'READ_ERROR', 'FAIL', None, None, None, None, None, None, None],
    ['X06', 'READ_ERROR', 'FAIL', None, None, None, None, None, None, None],
    ['N01', 'NO_SCHEMA', 'FAIL', None, None, None, None, None, None, None],
]
LOAD_COLS = ['Case', 'TargetType', 'SourceIDs', 'ExpOutcome', 'ExpReason', 'ExpRows', 'ExpErrorCells', 'Note']
LOAD = [
    ['L01', 'TestCsv', 'C01', 'OK', None, 5, 0, 'Еталон'],
    ['L02', 'TestCsv', 'C02', 'OK', None, 5, 0, 'Службові рядки відсічено'],
    ['L03', 'TestCsv', 'C03', 'OK', None, 10, 2, 'Report: помилки в клітинках'],
    ['L04', 'TestCsv', 'C04', 'ERROR', 'ПеревищенийПоріг', None, None, None],
    ['L05', 'TestCsv', 'C05', 'ERROR', 'ПомилкиТипізації', None, None, None],
    ['L06', 'TestCsv', 'C06', 'OK', None, 10, 0, 'Lenient: помилки стали null'],
    ['L07', 'TestCsv', 'C07', 'ERROR', 'ВідсутніОбов\'язковіПоля', None, None, None],
    ['L08', 'TestCsv', 'C08', 'OK', None, 5, 0, None],
    ['L09', 'TestCsv', 'C09', 'OK', None, 5, 0, 'Дані з першої колонки Назва'],
    ['L10', 'TestCsv', 'C10', 'OK', None, 0, None, None],
    ['L11', 'TestCsv', 'C11', 'OK', None, 0, None, 'Необов\'язковий відсутній'],
    ['L12', 'TestCsv', 'C12', 'ERROR', 'ФайлНедоступний', None, None, None],
    ['L13', 'TestCsv', 'C13', 'ERROR', 'ВідсутніОбов\'язковіПоля', None, None, None],
    ['L14', 'TestCsv', 'C14', 'OK', None, 5, 0, None],
    ['L15', 'TestCsv', 'C15', 'OK', None, 0, None, 'Вимкнене: даних немає'],
    ['L16', 'TestCsv', 'C01,C02', 'OK', None, 10, 0, 'Об\'єднання двох джерел'],
    ['L17', 'TestCsv', 'C01,C16', 'ERROR', 'ДублюванняПеріоду', None, None, None],
    ['L18', 'TestCsv', 'C01,C17', 'ERROR', 'ДублюванняФайлу', None, None, None],
    ['L19', 'TestXlsx', 'X01', 'OK', None, 5, 0, None],
    ['L20', 'TestXlsx', 'X02', 'OK', None, 5, 1, '#N/A лишається помилкою (Report)'],
    ['L21', 'TestXlsx', 'X03', 'ERROR', 'ЛистАбоТаблицюНеЗнайдено', None, None, None],
    ['L22', 'TestXlsx', 'X04', 'ERROR', 'НеоднозначнийВибір', None, None, None],
    ['L23', 'TestXlsx', 'X05', 'ERROR', None, None, None, 'Причину не перевіряємо'],
    ['L24', 'TestXlsx', 'X06', 'ERROR', 'НепідтримуванийФормат', None, None, None],
    ['L25', 'TestNoSchema', 'N01', 'ERROR', 'СхемаНеЗнайдена', None, None, None],
]

wb = Workbook(); ws = wb.active; ws.title = 'Інструкція'
for line in open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'README_stageC.md'), encoding='utf-8').read().splitlines():
    ws.append([line])
ws.column_dimensions['A'].width = 140

def plain(title, header, rows, widths=None):
    s = wb.create_sheet(title); s.append(header)
    for r in rows: s.append(r)
    for i, h in enumerate(header): s.column_dimensions[chr(65 + i)].width = max(12, len(h) + 4)
    return s
def table(title, name, header, rows):
    s = plain(title, header, rows)
    t = Table(displayName=name, ref=f'A1:{chr(64 + len(header))}{len(rows) + 1}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); s.add_table(t)
    return s

REG_HDR = ['SourceID', 'TargetType', 'FullPath', 'Path', 'FileName', 'TableSheetName', 'PeriodType', 'PeriodNumber',
           'IsRequired', 'Enabled', 'ErrorPolicy', 'ErrorThreshold', 'Опис (не копіювати)']
s = plain('Реєстр', REG_HDR, [[r[0], r[1], None, None, r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10]] for r in REG])
for row in s.iter_rows(min_row=2, min_col=12, max_col=12):
    for c in row: c.number_format = '0%'
plain('Схема', ['TargetType', 'SourceColumnName', 'TargetColumnName', 'DataType', 'IsRequired', 'TrimFlag', 'CustomTransform', 'Use'], SCHEMA)
plain('Фільтри', ['TargetType', 'SourceColumnName', 'Operator', 'Value', 'Use'], FILTERS)
plain('Призначення', ['TargetType', 'Опис'], TARGETS)
table('Очікування', 'tblTestExpected', EXP_COLS, EXP)
table('Навантаження', 'tblTestLoad', LOAD_COLS, LOAD)
wb.save(os.path.join(OUT, 'test_config.xlsx'))
print('config ok:', len(REG), 'sources,', len(EXP), 'summary cases,', len(LOAD), 'load cases')
